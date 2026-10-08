"""Builds sample.pptx covering the cases the converter handles. Needs: pip install python-pptx pillow"""
import io, sys
from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from PIL import Image

out = sys.argv[1] if len(sys.argv) > 1 else 'sample.pptx'
prs = Presentation()
prs.core_properties.title = 'Intro to Study Skills'

def png(color):
    buf = io.BytesIO(); Image.new('RGB', (320, 200), color).save(buf, 'PNG'); buf.seek(0); return buf

# 1 title slide
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = 'Intro to Study Skills'
s.placeholders[1].text = 'Spring semester'

# 2 bullets: nested, bold, italic, link; has notes
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = 'Plan your week'
tf = s.placeholders[1].text_frame
tf.text = 'Block out study time'
p = tf.add_paragraph(); p.text = 'Morning is best'; p.level = 1
p = tf.add_paragraph()
r = p.add_run(); r.text = 'Use a '; 
r = p.add_run(); r.text = 'planner'; r.font.bold = True
r = p.add_run(); r.text = ' & <calendar> '; r.font.italic = True
r = p.add_run(); r.text = 'guide'; r.hyperlink.address = 'https://example.edu/guide'
p = tf.add_paragraph(); r = p.add_run(); r.text = 'bad link'; r.hyperlink.address = 'javascript:alert(1)'
s.notes_slide.notes_text_frame.text = 'Remind students to bring their planners.'

# 3 picture with alt text + text -> side by side
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = 'Where to study'
pic = s.shapes.add_picture(png('steelblue'), Inches(1), Inches(2), Inches(3))
pic._element.nvPicPr.cNvPr.set('descr', 'A blue library desk')
tb = s.shapes.add_textbox(Inches(5), Inches(2), Inches(4), Inches(2)); tb.text_frame.text = 'Quiet places help focus.'

# 4 table
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = 'Weekly schedule'
t = s.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
for r_, row in enumerate([('Day', 'Task'), ('Mon', 'Read ch. 1'), ('Tue', 'Quiz')]):
    for c_, v in enumerate(row): t.cell(r_, c_).text = v

# 5 hidden slide
s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = 'Hidden slide'
s._element.set('show', '0')

# 6 no title, text box + image without alt
s = prs.slides.add_slide(prs.slide_layouts[6])
tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)); tb.text_frame.text = 'Untitled slide text'
s.shapes.add_picture(png('darkorange'), Inches(1), Inches(3), Inches(2))
s.shapes.add_picture(png('seagreen'), Inches(4), Inches(3), Inches(2))

# 7 chart (not convertible) + title only
s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = 'Results'
cd = CategoryChartData(); cd.categories = ['A', 'B']; cd.add_series('S', (1, 2))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(5), Inches(3), cd)

prs.save(out)
print('wrote', out)
